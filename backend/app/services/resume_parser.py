from io import BytesIO

import pdfplumber
from docx import Document


def _extract_pdf(content: bytes) -> str:
    text_parts = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text(filename: str, content: bytes) -> str:
    """Достаёт текст резюме из PDF/DOCX/TXT. Бросает ValueError на неподдерживаемых форматах."""
    if not filename or "." not in filename:
        raise ValueError("Не удалось определить формат файла")

    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        return _extract_pdf(content)
    if ext == "docx":
        return _extract_docx(content)
    if ext == "txt":
        return content.decode("utf-8", errors="ignore")

    raise ValueError(
        f"Формат .{ext} не поддерживается. Загрузите PDF, DOCX или TXT"
    )
